import os
import json
import tempfile
import mammoth
import anthropic
from flask import Flask, request, render_template, send_file, jsonify
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re

app = Flask(__name__)
app.config['MAX_CONTENT_LENGTH'] = 50 * 1024 * 1024  # 50MB max upload

ANTHROPIC_API_KEY = os.environ.get("ANTHROPIC_API_KEY")

SYSTEM_PROMPT = """You are a professional municipal code editor performing a non-substantive clerical review. 
Your role is to identify and correct only clerical, formatting, and drafting errors — never substantive legal or policy changes.

Review ONLY for:
- Spelling, grammar, and punctuation errors
- Numbering, lettering, and outline consistency
- Internal cross-references that appear broken, incorrect, or missing
- Inconsistent terminology, defined terms, or capitalization
- Missing words, repeated words, duplicated text, or misplaced text
- Formatting issues that may affect readability or interpretation

For each issue found, you MUST return:
1. The exact section number from the text
2. The issue type
3. The EXACT verbatim problematic text (no paraphrasing)
4. A recommended correction
5. Confidence level: High, Medium, or Possible issue

IMPORTANT RULES:
- Quote problematic text EXACTLY as it appears in the document
- Do not make substantive legal changes
- If a section number is unclear, write: "Section number not clear from provided file"
- If a cross-reference cannot be verified, note: "Unable to verify from provided file"
- Flag items needing drafting authority as High confidence but note they require verification

You MUST respond with valid JSON only, in exactly this structure:
{
  "summary": "Brief assessment of overall drafting quality and error categories found.",
  "total_issues": 0,
  "high_confidence": 0,
  "medium_confidence": 0,
  "possible_issues": 0,
  "requires_authority": ["list of section numbers requiring drafting authority determination"],
  "issues": [
    {
      "location": "§ X-XXX",
      "issue_type": "Type of error",
      "problematic_text": "Exact verbatim quote from the document",
      "recommended_correction": "Specific correction",
      "confidence": "High"
    }
  ],
  "cross_references_checked": [
    {
      "citation": "§ X-XXX → § Y-YYY",
      "finding": "Description of what was found",
      "status": "Verified" or "Wrong" or "Unable to verify"
    }
  ],
  "clean_draft_items": [
    {
      "section": "§ X-XXX",
      "change_from": "original text",
      "change_to": "corrected text",
      "note": "optional note if item cannot be redrafted"
    }
  ]
}"""


def extract_text_from_docx(file_path):
    """Extract plain text from a .docx file using mammoth."""
    with open(file_path, "rb") as f:
        result = mammoth.extract_raw_text(f)
    return result.value


def run_claude_review(text):
    """Send document text to Claude API and get structured review."""
    client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)

    # Truncate if very long (Claude's context limit)
    max_chars = 180000
    if len(text) > max_chars:
        text = text[:max_chars] + "\n\n[Document truncated due to length — review covers text above only]"

    message = client.messages.create(
        model="claude-sonnet-4-20250514",
        max_tokens=8000,
        system=SYSTEM_PROMPT,
        messages=[
            {
                "role": "user",
                "content": f"""Please perform a complete clerical review of the following municipal code document. 
Follow the required workflow exactly:
1. Read the full text and flag all potential issues
2. Verify each flagged issue against surrounding text
3. Check at least 20% of internal cross-references
4. Return your complete findings as valid JSON only.

DOCUMENT TEXT:
{text}"""
            }
        ]
    )

    raw = message.content[0].text.strip()
    # Strip markdown code fences if present
    raw = re.sub(r'^```json\s*', '', raw)
    raw = re.sub(r'\s*```$', '', raw)
    return json.loads(raw)


def set_cell_shading(cell, fill_color):
    """Apply background shading to a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_color)
    tcPr.append(shd)


def set_cell_borders(cell):
    """Set thin borders on a table cell."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), '4')
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), 'CCCCCC')
        tcBorders.append(element)
    tcPr.append(tcBorders)


def build_report_docx(review_data, output_path):
    """Generate a formatted .docx report from Claude's review JSON."""
    doc = Document()

    # Page margins
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    # ── Title block ──────────────────────────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("MUNICIPAL CODE CLERICAL REVIEW")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run("Non-Substantive Error Report — AI-Assisted Clerical Review").italic = True

    doc.add_paragraph()

    # ── Summary ───────────────────────────────────────────────────────────────
    h = doc.add_paragraph()
    h.add_run("SUMMARY AND ASSESSMENT").bold = True
    h.runs[0].font.size = Pt(13)
    h.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    doc.add_paragraph(review_data.get("summary", ""))

    stats = doc.add_paragraph()
    total = review_data.get("total_issues", 0)
    high = review_data.get("high_confidence", 0)
    med = review_data.get("medium_confidence", 0)
    poss = review_data.get("possible_issues", 0)
    stats.add_run(f"Total issues identified: {total}  |  High confidence: {high}  |  Medium: {med}  |  Possible: {poss}")
    stats.runs[0].bold = True

    requires = review_data.get("requires_authority", [])
    if requires:
        doc.add_paragraph(f"Items requiring drafting authority determination: {', '.join(requires)}")

    doc.add_paragraph()

    # ── Issues Table ──────────────────────────────────────────────────────────
    h2 = doc.add_paragraph()
    h2.add_run("ISSUES TABLE").bold = True
    h2.runs[0].font.size = Pt(13)
    h2.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    issues = review_data.get("issues", [])
    if issues:
        cols = 5
        col_widths = [Inches(0.9), Inches(1.3), Inches(2.1), Inches(2.1), Inches(0.85)]
        headers = ["Location", "Issue Type", "Problematic Text (exact quote)", "Recommended Correction", "Confidence"]
        header_color = "1F497D"
        alt_color = "EBF3FB"

        table = doc.add_table(rows=1, cols=cols)
        table.style = 'Table Grid'

        # Header row
        hrow = table.rows[0]
        for i, (cell, width, hdr) in enumerate(zip(hrow.cells, col_widths, headers)):
            cell.width = width
            set_cell_shading(cell, header_color)
            p = cell.paragraphs[0]
            run = p.add_run(hdr)
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(9)

        # Data rows
        for idx, issue in enumerate(issues):
            row = table.add_row()
            values = [
                issue.get("location", ""),
                issue.get("issue_type", ""),
                issue.get("problematic_text", ""),
                issue.get("recommended_correction", ""),
                issue.get("confidence", "")
            ]
            fill = alt_color if idx % 2 == 0 else "FFFFFF"
            for cell, width, val in zip(row.cells, col_widths, values):
                cell.width = width
                set_cell_shading(cell, fill)
                set_cell_borders(cell)
                p = cell.paragraphs[0]
                r = p.add_run(str(val))
                r.font.size = Pt(8.5)
                # Highlight High confidence in red, Medium in orange
                conf = issue.get("confidence", "")
                if cell == row.cells[4]:
                    if conf == "High":
                        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
                        r.bold = True
                    elif conf == "Medium":
                        r.font.color.rgb = RGBColor(0xE3, 0x6C, 0x09)
    else:
        doc.add_paragraph("No issues identified.")

    doc.add_paragraph()

    # ── Cross-Reference Spot-Check ────────────────────────────────────────────
    h3 = doc.add_paragraph()
    h3.add_run("CROSS-REFERENCE SPOT-CHECK").bold = True
    h3.runs[0].font.size = Pt(13)
    h3.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    xrefs = review_data.get("cross_references_checked", [])
    if xrefs:
        xr_cols = 3
        xr_widths = [Inches(1.8), Inches(4.5), Inches(1.4)]
        xr_headers = ["Citation in Text", "Section Referenced / Finding", "Status"]

        xtable = doc.add_table(rows=1, cols=xr_cols)
        xtable.style = 'Table Grid'

        xhrow = xtable.rows[0]
        for cell, width, hdr in zip(xhrow.cells, xr_widths, xr_headers):
            cell.width = width
            set_cell_shading(cell, header_color)
            p = cell.paragraphs[0]
            run = p.add_run(hdr)
            run.bold = True
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            run.font.size = Pt(9)

        for idx, xref in enumerate(xrefs):
            row = xtable.add_row()
            status = xref.get("status", "")
            values = [
                xref.get("citation", ""),
                xref.get("finding", ""),
                status
            ]
            fill = alt_color if idx % 2 == 0 else "FFFFFF"
            for i, (cell, width, val) in enumerate(zip(row.cells, xr_widths, values)):
                cell.width = width
                set_cell_shading(cell, fill)
                set_cell_borders(cell)
                p = cell.paragraphs[0]
                r = p.add_run(str(val))
                r.font.size = Pt(8.5)
                if i == 2:
                    if "Wrong" in status:
                        r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
                        r.bold = True
                    elif "Verified" in status:
                        r.font.color.rgb = RGBColor(0x37, 0x86, 0x10)
                        r.bold = True
    else:
        doc.add_paragraph("No cross-references checked.")

    doc.add_paragraph()

    # ── Clean Draft ───────────────────────────────────────────────────────────
    h4 = doc.add_paragraph()
    h4.add_run("CLEAN DRAFT — CORRECTED TEXT").bold = True
    h4.runs[0].font.size = Pt(13)
    h4.runs[0].font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    intro = doc.add_paragraph(
        "Corrected text is provided below for clearly non-substantive changes only. "
        "Sections with unfilled placeholders, substantive ambiguities, or items requiring "
        "legal or drafting authority verification are noted but not redrafted."
    )

    clean = review_data.get("clean_draft_items", [])
    for item in clean:
        sec = item.get("section", "")
        frm = item.get("change_from", "")
        to = item.get("change_to", "")
        note = item.get("note", "")

        sec_para = doc.add_paragraph()
        sec_para.add_run(f"{sec}").bold = True

        if frm and to:
            chg = doc.add_paragraph()
            chg.add_run("Change: ").bold = True
            chg.add_run(f"'{frm}'")
            chg2 = doc.add_paragraph()
            chg2.add_run("To: ").bold = True
            chg2.add_run(f"'{to}'")
        if note:
            n = doc.add_paragraph()
            n.add_run("Note: ").bold = True
            n.add_run(note).italic = True

        doc.add_paragraph()

    doc.save(output_path)


@app.route("/")
def index():
    return render_template("index.html")


@app.route("/review", methods=["POST"])
def review():
    if not ANTHROPIC_API_KEY:
        return jsonify({"error": "ANTHROPIC_API_KEY not configured on server"}), 500

    pasted_text = request.form.get("text", "").strip()
    uploaded_file = request.files.get("file")

    if pasted_text:
        text = pasted_text
    elif uploaded_file:
        if not uploaded_file.filename.endswith(".docx"):
            return jsonify({"error": "Please upload a .docx file"}), 400
        with tempfile.TemporaryDirectory() as tmpdir:
            input_path = os.path.join(tmpdir, "input.docx")
            uploaded_file.save(input_path)
            text = extract_text_from_docx(input_path)
    else:
        return jsonify({"error": "Please upload a .docx file or paste text"}), 400

    if not text.strip():
        return jsonify({"error": "Could not extract text from document"}), 400

    review_data = run_claude_review(text)

    with tempfile.TemporaryDirectory() as tmpdir:
        output_path = os.path.join(tmpdir, "review_report.docx")
        build_report_docx(review_data, output_path)
        return send_file(
            output_path,
            as_attachment=True,
            download_name="Municipal_Code_Clerical_Review.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    app.run(host="0.0.0.0", port=port)
